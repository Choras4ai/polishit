#!/usr/bin/env python3

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
import os


VERSION = 'V1.6.2'
DATE_TEXT = '2026年5月'
AUTHOR = '陈实之'
PRODUCT_NAME = '润石 PoliShit'
SITE_URL = 'https://www.runshi.top'
REPO_URL = 'https://github.com/Choras4ai/polishit'
RELEASE_URL = 'https://github.com/Choras4ai/polishit/releases/tag/v1.6.2'


doc = Document()


def setup_page():
    for section in doc.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(2.7)
        section.right_margin = Cm(2.7)


def setup_styles():
    style = doc.styles['Normal']
    style.font.name = '宋体'
    style.font.size = Pt(11)
    style.paragraph_format.line_spacing = 1.6
    style.paragraph_format.first_line_indent = Pt(22)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

    for level in range(1, 4):
        heading = doc.styles[f'Heading {level}']
        heading.font.name = '黑体'
        heading.font.color.rgb = RGBColor(0, 0, 0)
        heading.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
        heading.paragraph_format.first_line_indent = Pt(0)
        if level == 1:
            heading.font.size = Pt(18)
        elif level == 2:
            heading.font.size = Pt(14)
        else:
            heading.font.size = Pt(12)


def add_para(text):
    doc.add_paragraph(text)


def add_cover():
    for _ in range(5):
        doc.add_paragraph()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.first_line_indent = Pt(0)
    run = title.add_run('PCG 校园 AI 产品创意大赛')
    run.font.size = Pt(26)
    run.bold = True
    run.font.name = '黑体'
    run.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.first_line_indent = Pt(0)
    run = subtitle.add_run('初赛说明文档')
    run.font.size = Pt(22)
    run.font.name = '黑体'
    run.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')

    doc.add_paragraph()

    product = doc.add_paragraph()
    product.alignment = WD_ALIGN_PARAGRAPH.CENTER
    product.paragraph_format.first_line_indent = Pt(0)
    run = product.add_run(PRODUCT_NAME)
    run.font.size = Pt(20)
    run.bold = True
    run.font.name = '黑体'
    run.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')

    intro = doc.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    intro.paragraph_format.first_line_indent = Pt(0)
    intro.add_run('一个长在写作流程里的中文 AI 助手').font.size = Pt(13)

    doc.add_paragraph()

    for line in [
        f'版本：{VERSION}',
        f'日期：{DATE_TEXT}',
        f'作者：{AUTHOR}',
        f'官网：{SITE_URL}',
        f'仓库：{REPO_URL}',
        f'下载：{RELEASE_URL}',
    ]:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.add_run(line).font.size = Pt(11)

    doc.add_page_break()


def add_sections():
    doc.add_heading('一、为什么做这件事', level=1)
    add_para(
        '我自己是一个常年要写中文东西的人，论文、周报、推文、邮件，几乎每天都在写。最近一两年大模型已经足够好用了，'
        '但真正落到我的写作流程里时，体验并没有想象中那么顺。每写一段话，我都要先把它复制到聊天窗口，'
        '再切回原来的文档对照修改，遇到不满意的地方还要继续追问。模型本身不是问题，问题是它待在了一个跟写作过程脱节的窗口里。'
    )
    add_para(
        '我做润石 PoliShit 的初衷，就是想把这件事顺过来。让 AI 不再是“另一个网页”，而是一个会读我当前选区、'
        '会按我的方式给出修改建议、并且改完能直接落回原文档的桌面工具。它的目标场景很朴素：'
        '我在 Word、WPS、浏览器、飞书、邮件里写中文的时候，能一键润色，一键去掉那种一眼就能看出是 AI 写的腔调。'
    )

    doc.add_heading('二、润石到底是一个什么产品', level=1)
    add_para(
        '润石是一款跨 macOS 和 Windows 的中文写作辅助桌面应用，用 Electron 实现，目前发布到了 1.6.2。'
        '它的核心交互不是“打开聊天窗口和模型说话”，而是“在任何地方选中一段中文，让 AI 帮我改得更好”。'
    )
    add_para(
        '产品当前主要提供两件事：第一是润色，把句子里的语病、错别字、标点和别扭的搭配修掉，'
        '尽量不去改我原本想表达的意思；第二是降 AIGC，把那种一看就是模型写的、模板感很重的句子重新调成更像普通人写的样子。'
        '不管是哪一种，它给的都不是一段“最终稿”，而是逐条修改建议，可以一条条接受，也可以整篇替换，改错了还能撤销。'
    )
    add_para(
        '它和聊天产品最大的区别在于位置。聊天产品要求用户走过去、把内容贴进去、再贴回来。'
        '润石把这条路反了过来，它走到用户当前正在写作的窗口附近，由产品来配合写作流程，而不是反过来。'
    )

    doc.add_heading('三、它解决了什么具体问题', level=1)
    add_para(
        '过去我自己的中文润色流程大概是这样：写一段、不满意、Cmd+C、切到模型聊天页面、Cmd+V、'
        '加一句“帮我润色一下”、等结果出来、肉眼比对哪些动了、再 Cmd+C 贴回来、再读一遍是不是改坏了原意。'
        '这一整套动作下来，时间不长，但心情很碎。如果一篇稿子要改十几段，光是来回切换就足够把思路打散。'
    )
    add_para(
        '润石把这一段切碎的体验拼了回去。我只需要在原文里选中要改的句子，浮窗自动出现在选区附近，'
        '点一下润色或者降 AIGC，结果页里直接告诉我哪几个字被改了、为什么这么改，我可以一条条决定要不要接受。'
        '中间不需要打开任何额外的窗口，也不需要再去做手工 diff。整个过程从“去找 AI 帮忙”变成了“AI 在我旁边帮忙”。'
    )
    add_para(
        '这听上去只是少了几步操作，但实际写作时差别非常明显。它让“顺手改一下”这件事真的可以做到顺手，'
        '而不是每次都要下一个“值不值得为这一句话再切一次窗口”的决心。'
    )

    doc.add_heading('四、谁会用它', level=1)
    add_para(
        '从我接触到的用户来看，润石最自然的用户群是三类。第一类是在写论文、写作业、写邮件的学生和研究者，'
        '他们对“保留原意但表达更自然”这件事特别敏感，因为这直接关系到能不能交差、能不能被导师认可。'
        '第二类是日常需要写汇报、方案、对外邮件的职场用户，他们的核心诉求是把草稿快速打磨成可以发出去的正式文本。'
        '第三类是已经习惯先用 AI 起草、再人工精修的内容创作者，他们最受不了的是模型写出来那股“一眼 AI”的味道，'
        '降 AIGC 这件事对他们的价值会比润色还高。'
    )
    add_para(
        '这三类用户有一个共同点：写作不是他们的主业，但写作是他们绕不过去的一环。他们没有耐心学习一个复杂的新工具，'
        '只是希望在自己已经熟悉的环境里，多一个干活更顺的小帮手。这正好对应了润石的设计取向：尽量隐形、尽量轻、尽量不打断。'
    )

    doc.add_heading('五、它是怎么工作的', level=1)
    add_para(
        '在用户的视角里，润石的工作过程很短：选中文字，浮窗出现，点击或按快捷键，结果页出现，逐条接受，回到原文继续写。'
        '但在实现上，这条短路径背后串了好几块组件协作。'
    )
    add_para(
        '主进程负责整个生命周期，监听全局快捷键和选区状态，决定什么时候弹浮窗、什么时候走快捷键直接进入处理。'
        '选区监听模块在 macOS 上读取系统选区信息，在 Windows 上则通过复制回退的方式拿到当前选中的内容，'
        '这一段处理是产品体验最关键的环节，决定了“我以为它能拿到我选的那段话”到底是不是真的成立。'
    )
    add_para(
        'AI 处理这一层不是简单地把文本扔给一个模型，而是走过一个叫 AgentPipeline 的多步流程。'
        '以润色为例，它会先做语法层的纠错，再叠一层风格调整；以降 AIGC 为例，它会先做主体改写，再做一次更偏自然语感的修整。'
        '这一层的存在让结果更稳定，也让我可以针对不同模式调不同提示词，而不需要把所有要求堆给一个调用。'
    )
    add_para(
        '模型这一侧通过 Provider Factory 统一接入，目前已经支持自配的主流大模型 API，也支持调用本地的 Ollama 模型，'
        '当然也可以选择走润石自己的会员托管。三种路径的差异主要在“数据从哪里走、谁来承担成本、出问题找谁”，用户可以根据自己的场景挑。'
    )
    add_para(
        '结果出来之后，差异引擎会算出原文和新文之间的字符级差别，再以一条条修改建议的形式展示出来。'
        '用户可以一条条接受、忽略，也可以直接整篇替换。如果是浮窗触发的场景，接受之后会原位把改动写回原应用，'
        '中间出了任何问题，撤销入口可以把上一步的状态拉回来。'
    )

    doc.add_heading('六、为什么说它是 AI Native，不是套壳', level=1)
    add_para(
        '“AI Native”这个词最近被用得有点泛，我想说清楚我自己理解的边界。润石不是把一个聊天框塞到桌面上、'
        '换个图标就开张。它从输入到输出都是围绕模型能力重新设计的：输入端不要求用户做任何额外的整理，直接读选区；'
        '处理端不止一次模型调用，而是按任务拆成多步；输出端不交一份“最终答案”，而是把改动还原成可审阅的差异。'
    )
    add_para(
        '换句话说，离开了大模型，这个产品的所有交互都不成立。它的存在前提就是“现在的模型已经能做有意义的中文润色”，'
        '它的产品价值则在于把这件事用一种不打扰用户的方式落到真实写作环境里。这就是我理解的 AI Native：'
        '不是因为接了 AI 才叫 AI Native，而是因为没有 AI 就没有产品。'
    )

    doc.add_heading('七、目前做到了哪一步', level=1)
    add_para(
        '到比赛初赛这个时间点，润石已经不是一个只跑得起来的原型，它是一款已经发到 GitHub Releases、'
        '可以被陌生人下载安装、能在真实写作场景里用起来的桌面应用。当前版本是 1.6.2，macOS 和 Windows 都有对应安装包。'
    )
    add_para(
        '产品侧，浮窗触发、全局快捷键、润色与降 AIGC 两个核心模式、结果页逐条审阅、原位写回、撤销，这一整套主流程是通的。'
        '设置页可以选模型、配 API、配自定义提示词，引导页会带新用户走完一遍最小可用的路径。'
        '官网上线在 runshi.top，承担产品介绍、版本说明、隐私协议、使用条款和下载入口。'
    )
    add_para(
        '后端侧，会员托管、积分扣减、模型分层、订单流程、自动化更新这些基础能力都已经实现并跑通，'
        '关键链路有自动化测试覆盖。支付这一块代码已经接好支付宝和微信的形态，只是真实商户联调还没完成，所以目前不对外宣称已经可付费使用。'
    )
    add_para(
        '我并不想把完成度说得比实际更满。今天它的状态是：核心写作体验已经可以稳定工作，商业闭环只差最后一公里，'
        '团队协作和更高级的写作模板还在规划里。但作为一个由我一个人完成、并且已经发布到公网的产品，我认为它具备了继续往真实用户里走的资格。'
    )
    add_para(
        '关于“最后一公里”这件事，可以多说几句背景。润石要让用户真正用起来，绕不开两件很现实的事：一是公网备案，'
        '二是支付。这两件事在国内都有一个硬门槛，必须以个体工商户或者公司的身份去办，个人开发者身份是走不通的。'
        '我之前为此专门注册了一个个体工商户，本来想用这个主体去申请武汉光谷面向创业者的算力资源，'
        '希望把模型推理这一块的成本压一压，让润石可以更从容地给学生和轻量用户提供免费额度。'
        '可惜那一次申请最终没有通过，整个流程走下来时间也不短，对项目节奏其实是一次不小的影响。'
    )
    add_para(
        '现在我重新调整了思路，正在准备以公司主体重新注册，把备案和支付这两件事彻底推到正式上线状态。'
        '走公司这条路虽然比个体户重一些，但它可以同时把对公账户、发票、支付通道、合规备案这些后续一定会遇到的问题一次性铺好，'
        '也让润石在面对未来可能的渠道合作和企业用户时不会再因为主体资质卡住。'
        '换句话说，今天产品本身的能力已经到位，剩下的更多是行政和资质层面的工作，我会在比赛之后继续把它推完。'
    )

    doc.add_heading('八、它和市面上其他工具有什么不一样', level=1)
    add_para(
        '通用聊天产品胜在通用，但用在中文写作场景里有两个明显的不顺：一个是位置不对，要求用户主动去找它；'
        '另一个是结果形态不对，习惯整段重写，让人很难判断它到底改了什么、为什么改。润石做的事正好是反着来：'
        '我去找你写作的位置，我把改动一条条摊给你看，决定权留在你手里。'
    )
    add_para(
        '传统语法检查工具的问题是另一种，它们大多基于规则，对中文语义层面的判断比较弱，遇到“改得对但读起来奇怪”的句子就显得无能为力。'
        '润石用模型来理解中文表达，再用产品交互来约束模型的输出形态，目标是把“懂中文”和“可控制”两件事同时拿下。'
    )
    add_para(
        '所以润石不去和谁比谁的聊天体验更好，也不去和谁比谁的规则更全。它的差异在于工作流：让中文写作里那段“反复打磨”的过程，'
        '可以在不离开当前文档的前提下完成。这是一件听起来不大、但用过之后就回不去的事。'
    )

    doc.add_heading('九、商业化是怎么想的', level=1)
    add_para(
        '商业化这件事我没有走纯订阅，原因很简单，目标用户里有一大部分是学生和轻度创作者，他们对“每个月固定花钱”这件事天然警惕，'
        '反而对“按用量买积分”更能接受。所以润石的方式是：免费试用先把基础能力体验完，重度场景用积分包覆盖，'
        '不同模型按能力档位设不同倍率，让用户自己去选“值不值”。'
    )
    add_para(
        '这种结构对成本侧也更友好。上游模型价格在变，不同模型适合不同场景，倍率机制可以把这些差异自然消化掉，'
        '而不需要每次涨价都去改套餐结构。同时它对用户也更透明：你这次花了多少积分、用的是哪一档模型，结果页里就能看到。'
    )
    add_para(
        '增长上，我目前优先做三件事：一是官网和 Releases 这条最朴素的下载链路，让感兴趣的人能一分钟内拿到安装包；'
        '二是用真实的“润色前 vs 润色后”对比内容做产品展示，因为这件事的价值不靠口播，要靠看；'
        '三是慢慢把使用过的人沉淀成口碑节点，特别是高校里写论文、写求职文书的场景。'
    )

    doc.add_heading('十、风险和边界', level=1)
    add_para(
        '我想很明确地说清楚润石不做什么。它是一个写作辅助工具，不是代写工具，更不承担任何“帮你绕过检测”的承诺。'
        '降 AIGC 功能的目标是减少模板化表达，让句子读起来更像人写的，而不是用来逃避任何特定检测系统。'
        '所有由模型产生的修改最终需要由用户自己审阅、决定是否采纳，由此带来的内容责任也由用户承担。'
    )
    add_para(
        '在数据路径上，自配 API 模式下文本直接发往用户自己填写的服务商，平台不经手用户的密钥；'
        '会员托管模式下文本经由润石后端转发，便于统一计费和监控；本地 Ollama 模式下文本可以完全留在用户本机，不出网。'
        '哪种模式更合适由用户自己选，产品在界面和文档里都明确标注了差异，不做任何模糊处理。'
    )
    add_para(
        '支付链路虽然代码已经准备好，但在真实商户联调完成之前，我不会对外宣称“已经可以正式付费使用”。'
        '这件事我宁可慢一点，也不能让用户在不清晰的状态下花钱。'
    )

    doc.add_heading('十一、Demo 与提交材料', level=1)
    add_para(
        f'产品官网在 {SITE_URL}，进入后可以看到产品介绍、版本说明和下载入口；'
        f'代码仓库托管在 {REPO_URL}，记录了完整开发历程；'
        f'安装包发布在 GitHub Releases，可以从 {RELEASE_URL} 直接下载 macOS 与 Windows 两个平台的安装包。'
        '说明文档就是这份 PDF，演示录屏会另附。'
    )
    add_para(
        '录屏我打算按真实使用过程来拍，而不是做成功能展示片。开头先复刻一下旧流程的痛点：在 Word 里写一段、'
        '复制到聊天窗口、贴回来、肉眼对比；然后切到润石的工作流：选中、浮窗、润色、逐条接受、改动直接写回原文档。'
        '中间会演示一次降 AIGC，对比改前改后的语感差异，再演示一次模型切换，让评委看到不同档位模型的实际效果区别。'
        '最后会展示一下会员模式下积分扣减的可视化界面，把整个商业闭环也走一遍。'
    )

    doc.add_heading('十二、接下来要做什么', level=1)
    add_para(
        '短期内我会优先把支付链路真正联调上，把当前已有功能的稳定性继续打磨，特别是 Windows 端选区识别在不同应用里的兼容性。'
        '中期会增加更多写作模板和风格规则，让润石不只是“通用润色”，而是可以记住用户偏好的写作助手。'
        '更长一些，会探索团队协作场景，比如团队统一文风、团队配额管理，把单人工具扩到组织里去。'
    )
    add_para(
        '比赛对我来说是一个很好的节点，它逼我把这一年的工作认真梳理一遍，看清自己做到了哪、哪里还差很多。'
        '不管最终结果如何，润石本身会作为一个长期项目继续走下去，因为这件事我自己在用，也确实从中受益。'
    )

    doc.add_heading('十三、写在最后', level=1)
    add_para(
        '我做润石的过程没有任何团队，所有的代码、产品决策、视觉、文案、官网、运营准备都是我一个人完成的。'
        '它不一定是最华丽的项目，但它是一个真的在跑、有人在用、并且我自己每天都在依赖的产品。'
        '我希望它代表的不是“又一个 AI 应用”，而是一种相对克制的产品观：'
        '不去包装新概念，不去假设用户会改变习惯，先把模型能力老老实实地装进用户已经在做的事情里，'
        '把那一段最让人烦的反复打磨变得轻一点。如果它能在更多人的写作过程里担起这个角色，那对我来说就值得了。'
    )


def save_documents():
    docs_dir = os.path.dirname(__file__)
    legacy_path = os.path.join(docs_dir, '润石PoliShit_产品介绍与技术白皮书_OPC认证.docx')
    contest_path = os.path.join(docs_dir, '润石PoliShit_PCG校园AI产品创意大赛_初赛说明文档.docx')
    doc.save(legacy_path)
    doc.save(contest_path)
    print(f'✅ 文档已生成: {legacy_path}')
    print(f'✅ 文档已生成: {contest_path}')


setup_page()
setup_styles()
add_cover()
add_sections()
save_documents()
